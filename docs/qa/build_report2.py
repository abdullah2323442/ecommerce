"""Builds the comprehensive QA report: every feature, what it is for, what was done, what was seen, and its own screenshots.

Inputs: docs/FEATURE-CHECKLIST.md, docs/qa/data/part*.txt (plain explanations), docs/qa/results/*.json (from the runner), docs/qa/shots2/*.jpg.
Output: docs/qa/Siams-Aqua-QA-Report-Final.docx
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).parent
SHOTS = ROOT / "shots2"
THUMBS = ROOT / "thumbs2"
THUMBS.mkdir(exist_ok=True)
RESULTS = ROOT / "results"
OUT = ROOT / "Siams-Aqua-QA-Report-Final.docx"
CHECKLIST = ROOT.parent / "FEATURE-CHECKLIST.md"

ROLE = {
    "admin": "the owner", "staff": "a staff member", "paikari": "a paikari shop", "paikari2": "a paikari shop", "paikari3": "a paikari shop",
    "wholesale": "a wholesaler", "visitor": "a signed-in shopper", "guest": "a stranger (not signed in)", "newguest": "a stranger (not signed in)",
    "vendor": "a restaurant owner", "mpo": "an MPO", "craftsman": "a craftsman",
}

LABEL = {
    "pass": ("Works", RGBColor(0x1B, 0x6B, 0x3A)),
    "fail": ("Does not work as described", RGBColor(0xB0, 0x1E, 0x1E)),
    "error": ("Could not be tested", RGBColor(0x9A, 0x6A, 0x00)),
    "skipped": ("Only partly shown", RGBColor(0x9A, 0x6A, 0x00)),
    "notbuilt": ("Not built yet", RGBColor(0x55, 0x55, 0x55)),
    "none": ("Not tested yet", RGBColor(0x55, 0x55, 0x55)),
}

# --- features
chapters, feats, current = [], {}, None
for line in CHECKLIST.read_text(encoding="utf-8").splitlines():
    if line.startswith("### "):
        current = line[4:].strip()
        chapters.append(current)
        continue
    m = re.match(r"\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*(Phase \d|Future)\s*\|\s*(\S+)\s*\|", line)
    if m and current:
        n = int(m.group(1))
        feats[n] = dict(n=n, name=m.group(2).replace("**", "").replace("\\*", "*"), chapter=current,
                        checklist={"✅": "Built", "🟡": "Partly built", "☐": "Not built yet"}.get(m.group(4), "Built"))
for part in sorted((ROOT / "data").glob("part*.txt")):
    for line in part.read_text(encoding="utf-8").splitlines():
        if line.strip():
            n, _b, _f, text = line.split("|", 3)
            feats[int(n)]["text"] = text.strip()
for n, f in feats.items():
    r = RESULTS / f"{n}.json"
    f["res"] = json.loads(r.read_text(encoding="utf-8")) if r.exists() else None
    f["status"] = f["res"]["status"] if f["res"] else ("notbuilt" if f["checklist"] == "Not built yet" else "none")


def thumb(name):
    src = SHOTS / name
    dst = THUMBS / name
    if not dst.exists() or dst.stat().st_mtime < src.stat().st_mtime:
        im = Image.open(src).convert("RGB")
        im.thumbnail((1000, 1000))
        im.save(dst, "JPEG", quality=60, optimize=True)
    return dst


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def small(par, size=8.5, color=None, italic=False, bold=False):
    for r in par.runs:
        r.font.size = Pt(size)
        r.italic = italic
        r.bold = bold
        if color:
            r.font.color.rgb = color


doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(1.6)
    s.top_margin = s.bottom_margin = Cm(1.6)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

doc.add_heading("Siam's Aqua: feature-by-feature QA report", level=0)
doc.add_paragraph(
    "Every feature in the project's checklist, tested as a real person would use it. For each one this report says what the feature is for, "
    "what the tester did, what the tester saw, and shows the screens along the way."
)
doc.add_heading("How to read it", level=2)
for t in (
    "Works: the tester used the feature and it did what it should.",
    "Does not work as described: the tester used the feature and it did not do what the checklist says. These are the findings to fix.",
    "Only partly shown: part of the feature was shown. The note says what could not be shown and why.",
    "Could not be tested: the test itself broke before reaching a verdict. It says where.",
    "Not built yet: the checklist says this is not built, so there is nothing to test.",
    "Not tested yet: this feature has not been reached in this pass. Its explanation is included so the list is complete.",
    "Everything was done on a copy of the development site with sample data, in English, as the person named in each picture.",
):
    doc.add_paragraph(t, style="List Bullet")

# at a glance
doc.add_heading("At a glance", level=2)
cols = ["pass", "fail", "skipped", "error", "notbuilt", "none"]
tbl = doc.add_table(rows=1, cols=2 + len(cols))
tbl.style = "Light Grid Accent 1"
for i, h in enumerate(["Area", "Features"] + [LABEL[c][0] for c in cols]):
    tbl.rows[0].cells[i].text = h
tot = [0] * (1 + len(cols))
for ch in chapters:
    fs = [f for f in feats.values() if f["chapter"] == ch]
    row = tbl.add_row().cells
    vals = [len(fs)] + [sum(f["status"] == c for f in fs) for c in cols]
    row[0].text = ch
    for i, v in enumerate(vals):
        row[i + 1].text = str(v)
        tot[i] += v
row = tbl.add_row().cells
row[0].text = "All features"
for i, v in enumerate(tot):
    row[i + 1].text = str(v)

# findings
bad = [f for f in feats.values() if f["status"] == "fail"]
if bad:
    doc.add_heading("Findings: what does not work as described", level=2)
    for f in sorted(bad, key=lambda x: x["n"]):
        misses = [c["text"] for c in f["res"]["checks"] if not c["ok"]]
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{f['n']}. {f['name']}: ").bold = True
        p.add_run("; ".join(misses[:3]) + (" ..." if len(misses) > 3 else ""))

for ch in chapters:
    doc.add_page_break()
    doc.add_heading(ch, level=1)
    for f in sorted((x for x in feats.values() if x["chapter"] == ch), key=lambda x: x["n"]):
        label, color = LABEL[f["status"]]
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(f"{f['n']}. {f['name']}")
        r.bold = True
        r.font.size = Pt(13)
        r2 = h.add_run(f"    {label}")
        r2.bold = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = color

        p = doc.add_paragraph()
        p.add_run("What it is for: ").bold = True
        p.add_run(f.get("text", ""))
        res = f["res"]
        if not res:
            if f["status"] == "notbuilt":
                q = doc.add_paragraph("The checklist says this is not built yet, so there was nothing to test.")
            else:
                q = doc.add_paragraph("Not tested yet in this pass.")
            small(q, 9, italic=True)
            continue
        if res.get("did"):
            p = doc.add_paragraph()
            p.add_run("What the tester did: ").bold = True
            p.add_run(res["did"])
        if res["checks"]:
            p = doc.add_paragraph()
            p.add_run("What the tester saw:").bold = True
            for c in res["checks"]:
                li = doc.add_paragraph(style="List Bullet")
                tag = li.add_run("Yes. " if c["ok"] else "No. ")
                tag.bold = True
                tag.font.color.rgb = LABEL["pass"][1] if c["ok"] else LABEL["fail"][1]
                li.add_run(c["text"])
        for note in res["notes"]:
            p = doc.add_paragraph()
            p.add_run("Note: ").bold = True
            p.add_run(note)
        if res.get("error"):
            p = doc.add_paragraph()
            p.add_run("Test stopped: ").bold = True
            p.add_run(res["error"])
        shots = res["shots"]
        for i in range(0, len(shots), 2):
            pair = shots[i:i + 2]
            t = doc.add_table(rows=2, cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j in range(2):
                cap, pic = t.rows[0].cells[j], t.rows[1].cells[j]
                if j < len(pair):
                    s = pair[j]
                    shade(cap, "EEF1F5")
                    cap.paragraphs[0].add_run(f"Step {i + j + 1}: {s['caption']}")
                    small(cap.paragraphs[0], 8.5, bold=True)
                    pic.paragraphs[0].add_run().add_picture(str(thumb(s["file"])), width=Cm(8.7))
                    c2 = pic.add_paragraph(f"Seen as {ROLE.get(s['role'], s['role'])}  |  {s['url']}")
                    small(c2, 7.5, RGBColor(0x66, 0x66, 0x66))

doc.save(OUT)
done = [f for f in feats.values() if f["res"]]
print(f"saved {OUT} ({OUT.stat().st_size / 1e6:.1f} MB): {len(feats)} features, {len(done)} tested")

"""Builds the full feature report (all 613 features) as one Word file.

Inputs:  docs/FEATURE-CHECKLIST.md (names, status, chapters), docs/qa/data/part*.txt (plain explanations and screen keys),
         docs/qa/shots/*.jpg (screens captured by the capture script).
Output:  docs/qa/Siams-Aqua-Feature-Report.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).parent
SHOTS = ROOT / "shots"
THUMBS = ROOT / "thumbs"
THUMBS.mkdir(exist_ok=True)
OUT = ROOT / "Siams-Aqua-Feature-Report.docx"
CHECKLIST = ROOT.parent / "FEATURE-CHECKLIST.md"
CAPTURE = ROOT / "capture.mjs"

ROLE_NAME = {
    "admin": "Owner's back office", "staff": "Staff screen", "paikari": "Paikari shop buyer", "wholesale": "Wholesale buyer/seller",
    "visitor": "Signed-in shopper", "guest": "Public site (not signed in)", "vendor": "Restaurant owner",
}

# --- screens: key -> (role, path) read from the capture script so the two never disagree
screens = {}
for m in re.finditer(r"\['(\w+)',\s*'(\w+)',\s*'([^']+)'\]", CAPTURE.read_text(encoding="utf-8")):
    screens[m.group(1)] = (m.group(2), m.group(3))

# --- checklist: chapters, names, status
chapters, feats = [], {}
current = None
for line in CHECKLIST.read_text(encoding="utf-8").splitlines():
    if line.startswith("### "):
        current = line[4:].strip()
        chapters.append(current)
        continue
    m = re.match(r"\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*(Phase \d|Future)\s*\|\s*(\S+)\s*\|", line)
    if m and current:
        n = int(m.group(1))
        name = m.group(2).replace("**", "").replace("\\*", "*")
        status = {"✅": "Built", "🟡": "Partly built", "☐": "Not built yet"}.get(m.group(4), "Built")
        feats[n] = dict(n=n, name=name, phase=m.group(3), status=status, chapter=current)

# --- plain explanations and screen keys
for part in sorted((ROOT / "data").glob("part*.txt")):
    for line in part.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        n, back, front, text = line.split("|", 3)
        feats[int(n)].update(back=None if back == "-" else back, front=None if front == "-" else front, text=text.strip())

missing = [n for n, f in feats.items() if "text" not in f]
assert not missing, f"features without an explanation: {missing}"


def thumb(key):
    src = SHOTS / f"{key}.jpg"
    if not src.exists():
        return None
    dst = THUMBS / f"{key}.jpg"
    if not dst.exists() or dst.stat().st_mtime < src.stat().st_mtime:
        im = Image.open(src).convert("RGB")
        im.thumbnail((900, 900))
        im.save(dst, "JPEG", quality=62, optimize=True)
    return dst


def shade(cell, hex_fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(1.6)
    s.top_margin = s.bottom_margin = Cm(1.6)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)

doc.add_heading("Siam's Aqua: every feature, explained and shown", level=0)
p = doc.add_paragraph(
    "This report goes through all 613 features in the project's feature checklist. For each one it says, in plain words, what the feature does "
    "for the business or the customer, and shows a screenshot of the screen where it lives: the owner's back office on the left, "
    "the customer or shop side on the right."
)
doc.add_heading("How to read it", level=2)
for t in (
    "Status (Built, Partly built, Not built yet) comes from the project's own feature checklist. It was not re-tested for this report.",
    "Screenshots were taken on the local development site with sample data, in English, as the role named under each picture.",
    "A screenshot shows the screen where the feature lives. Some features are rules that work behind the scenes, so there is no separate screen to show. Those say so.",
    "Features that are not built yet have no screenshot, because there is nothing to show.",
):
    doc.add_paragraph(t, style="List Bullet")

# summary table
doc.add_heading("At a glance", level=2)
tbl = doc.add_table(rows=1, cols=5)
tbl.style = "Light Grid Accent 1"
for i, h in enumerate(["Area", "Features", "Built", "Partly built", "Not built yet"]):
    tbl.rows[0].cells[i].text = h
tot = [0, 0, 0, 0]
for ch in chapters:
    fs = [f for f in feats.values() if f["chapter"] == ch]
    row = tbl.add_row().cells
    vals = [len(fs), sum(f["status"] == "Built" for f in fs), sum(f["status"] == "Partly built" for f in fs), sum(f["status"] == "Not built yet" for f in fs)]
    row[0].text = ch
    for i, v in enumerate(vals):
        row[i + 1].text = str(v)
        tot[i] += v
row = tbl.add_row().cells
row[0].text = "All features"
for i, v in enumerate(tot):
    row[i + 1].text = str(v)
    for r in row[i + 1].paragraphs[0].runs:
        r.bold = True

COL = {"Built": RGBColor(0x1B, 0x6B, 0x3A), "Partly built": RGBColor(0x9A, 0x6A, 0x00), "Not built yet": RGBColor(0x8A, 0x2A, 0x2A)}

for ch in chapters:
    doc.add_page_break()
    doc.add_heading(ch, level=1)
    for f in sorted((x for x in feats.values() if x["chapter"] == ch), key=lambda x: x["n"]):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(f"{f['n']}. {f['name']}")
        r.bold = True
        r.font.size = Pt(12)
        r2 = h.add_run(f"   {f['status']}")
        r2.bold = True
        r2.font.size = Pt(9)
        r2.font.color.rgb = COL[f["status"]]
        body = doc.add_paragraph(f["text"])
        body.paragraph_format.keep_with_next = True

        if f["status"] == "Not built yet" or not (f["back"] or f["front"]):
            note = doc.add_paragraph(
                "Nothing to show yet: this feature is not built." if f["status"] == "Not built yet"
                else "This works behind the scenes, so there is no separate screen to show."
            )
            note.runs[0].italic = True
            note.runs[0].font.size = Pt(9)
            continue

        t = doc.add_table(rows=2, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (label, key) in enumerate((("Back office", f["back"]), ("Customer side", f["front"]))):
            head, pic = t.rows[0].cells[i], t.rows[1].cells[i]
            shade(head, "EEF1F5")
            head.paragraphs[0].add_run(label).bold = True
            head.paragraphs[0].runs[0].font.size = Pt(8.5)
            path = thumb(key) if key else None
            if key and path:
                role, url = screens[key]
                pic.paragraphs[0].add_run().add_picture(str(path), width=Cm(8.4))
                cap = pic.add_paragraph(f"{ROLE_NAME[role]}  |  {url}")
                cap.runs[0].font.size = Pt(7.5)
                cap.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                text = "Screen not captured for this report (needs a restaurant owner login)." if key else "No separate screen on this side."
                pic.paragraphs[0].add_run(text).italic = True
                pic.paragraphs[0].runs[0].font.size = Pt(8.5)

doc.save(OUT)
print(f"saved {OUT} ({OUT.stat().st_size / 1e6:.1f} MB), {len(feats)} features, {len(screens)} screens")

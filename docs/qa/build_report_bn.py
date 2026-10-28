"""Builds the Bangla version of the QA report, for the client.

Inputs: the same checklist and results as build_report2.py, plus the Bangla text in docs/qa/bn/chNN.txt
Output: docs/qa/Siams-Aqua-QA-Report-Bangla.docx
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).parent
src = (ROOT / "build_report2.py").read_text(encoding="utf-8").split("def thumb")[0].replace("ROOT = Path(__file__).parent", "ROOT = Path('.').resolve()")
exec(src)  # chapters, feats (with res, status), SHOTS, THUMBS, ROOT

OUT = ROOT / "Siams-Aqua-QA-Report-Bangla.docx"
FONT = "Nirmala UI"

NAVY = "1F3A5F"
GREEN = "1B7A43"
GREEN_BG = "E3F4EA"
AMBER = "9A6A00"
AMBER_BG = "FFF3D6"
GREY_BG = "F2F4F7"
BLUE_BG = "E8EEF7"

ROLE_BN = {
    "admin": "মালিক", "staff": "একজন কর্মী", "paikari": "একটি পাইকারি দোকান", "paikari2": "একটি পাইকারি দোকান", "paikari3": "একটি পাইকারি দোকান",
    "wholesale": "একজন হোলসেলার", "wholesaler": "একজন হোলসেলার", "visitor": "একজন লগইন করা ক্রেতা", "buyer": "একজন ক্রেতা", "guest": "একজন অপরিচিত (লগইন করেননি)",
    "newguest": "একজন অপরিচিত (লগইন করেননি)", "vendor": "একজন রেস্টুরেন্ট মালিক", "mpo": "একজন এমপিও", "craftsman": "একজন কারিগর",
    "applicant": "একজন আবেদনকারী", "seller": "একজন বিক্রেতা", "rider": "একজন ডেলিভারি রাইডার", "driver": "একজন ডেলিভারি রাইডার", "courier": "একজন কুরিয়ার",
}

DIGITS = str.maketrans("0123456789", "০১২৩৪৫৬৭৮৯")


def bn(x):
    return str(x).translate(DIGITS)


def pct(a, b):
    return bn(f"{(100 * a / b):.1f}".rstrip("0").rstrip(".")) if b else "০"


# ------------------------------------------------------------------ Bangla text
BN = {}      # feature number -> dict(name, E, T, N, S[list])
CH_BN = {}   # chapter index (1-based) -> (title, intro)
for i, ch in enumerate(chapters, 1):
    path = ROOT / "bn" / f"ch{i:02d}.txt"
    if not path.exists():
        continue
    cur = None
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        if line.startswith("CH|"):
            _, title, intro = line.split("|", 2)
            CH_BN[i] = (title.strip(), intro.strip())
        elif line.startswith("@"):
            head, name = line[1:].split("|", 1)
            cur = BN[int(head)] = dict(name=name.strip(), E="", T="", N="", S=[])
        elif cur is not None:
            tag, _, rest = line.partition("|")
            if tag == "S":
                cur["S"].append(rest.strip())
            elif tag in ("E", "T", "N"):
                cur[tag] = rest.strip()

missing = [n for n, f in feats.items() if f["chapter"] in [chapters[i - 1] for i in CH_BN] and n not in BN]
if missing:
    print("MISSING Bangla for features:", missing)
    sys.exit(1)
for n, b in BN.items():
    f = feats[n]
    shots = f["res"]["shots"] if f["res"] else []
    if len(b["S"]) != len(shots):
        print(f"caption count differs for {n}: {len(b['S'])} vs {len(shots)}")
        sys.exit(1)
    if not b["E"] or not b["T"]:
        print("empty text for", n)
        sys.exit(1)
    if f["status"] == "skipped" and not b["N"]:
        print("a partly shown feature needs a note:", n)
        sys.exit(1)

use = [i for i in range(1, len(chapters) + 1) if i in CH_BN]
used_feats = [f for f in feats.values() if f["chapter"] in [chapters[i - 1] for i in use]]


def thumb(name):
    src_ = SHOTS / name
    dst = THUMBS / name
    if not dst.exists() or dst.stat().st_mtime < src_.stat().st_mtime:
        im = Image.open(src_).convert("RGB")
        im.thumbnail((1000, 1000))
        im.save(dst, "JPEG", quality=60, optimize=True)
    return dst


# ------------------------------------------------------------------ docx helpers
def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szcs)
    if bold is not None:
        run.bold = bold
        if bold:
            rpr.append(OxmlElement("w:bCs"))
    if italic:
        run.italic = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    for old in pr.findall(qn("w:shd")):
        pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def borders(cell, color="D5DAE1", sz=4, sides=("top", "left", "bottom", "right")):
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        if s in sides:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    pr.append(b)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    pr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    pr.append(m)


def width(cell, cm):
    cell.width = Cm(cm)
    tcpr = cell._tc.get_or_add_tcPr()
    w = tcpr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW")
        tcpr.append(w)
    w.set(qn("w:w"), str(int(cm * 567)))
    w.set(qn("w:type"), "dxa")


def fixed_layout(table):
    pr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    pr.append(lay)


def para(container, text="", size=11, bold=False, color=None, align=None, space_after=4, space_before=0, italic=False, keep=False):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.25
    if keep:
        p.paragraph_format.keep_with_next = True
    if align:
        p.alignment = align
    if text:
        set_font(p.add_run(text), size, bold, color, italic)
    return p


def first_para(cell, text, size=11, bold=False, color=None, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    if align:
        p.alignment = align
    set_font(p.add_run(text), size, bold, color)
    return p


def heading(doc, text, level):
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, {0: 28, 1: 20, 2: 15, 3: 12}[level], True, NAVY)
    return p


def add_field(par, instr, size=9, color="666666"):
    r = par.add_run()
    set_font(r, size, None, color)
    for typ, txt in (("begin", None), (None, instr), ("end", None)):
        if typ:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), typ)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        r._r.append(e)


# ------------------------------------------------------------------ the document
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = Cm(1.6)
sec.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]
st.font.name = FONT
st.element.rPr.rFonts.set(qn("w:cs"), FONT)
st.font.size = Pt(11)
for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
    s = doc.styles[name]
    s.font.name = FONT
    s.element.rPr.rFonts.set(qn("w:cs"), FONT)
    s.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    s.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)

lng = OxmlElement("w:lang")
lng.set(qn("w:val"), "bn-BD")
lng.set(qn("w:eastAsia"), "bn-BD")
lng.set(qn("w:bidi"), "bn-BD")
st.element.rPr.append(lng)
doc.core_properties.title = "সিয়ামস অ্যাকুয়া: ফিচার-ভিত্তিক পরীক্ষার প্রতিবেদন"
doc.core_properties.author = "সিয়ামস অ্যাকুয়া"
doc.core_properties.language = "bn-BD"

# footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("সিয়ামস অ্যাকুয়া: পরীক্ষার প্রতিবেদন   |   পাতা "), 9, None, "666666")
add_field(fp, "PAGE")

CONTENT_W = 17.4

# ---- cover
t = doc.add_table(rows=1, cols=1)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
fixed_layout(t)
c = t.rows[0].cells[0]
width(c, CONTENT_W)
shade(c, NAVY)
cell_margins(c, 500, 500, 400, 400)
first_para(c, "সিয়ামস অ্যাকুয়া", 30, True, "FFFFFF")
para(c, "প্রতিটি ফিচারের পরীক্ষার প্রতিবেদন", 18, True, "FFFFFF", space_after=8)
para(c, "আপনি যা যা চেয়েছিলেন, তার প্রতিটি আমরা একজন সত্যিকারের ব্যবহারকারীর মতো চালিয়ে দেখেছি। এই প্রতিবেদনে সহজ ভাষায় লেখা আছে প্রতিটি ফিচার কী কাজে লাগে, আমরা কী করে দেখেছি, আর কী দেখতে পেয়েছি। সাথে আছে পর্দার ছবি।", 11.5, False, "DCE6F2", space_after=10)
para(c, "তারিখ: " + bn("২৩ অক্টোবর ২০২৬") + "     |     প্রস্তুত করা হয়েছে: সিয়াম মাহবুবের জন্য", 10.5, False, "B9C8DD", space_after=0)

# ---- dashboard
heading(doc, "এক নজরে: পরীক্ষার ফল", 1)
total = len(used_feats)
works = sum(f["status"] == "pass" for f in used_feats)
partly = sum(f["status"] == "skipped" for f in used_feats)
broken = sum(f["status"] == "fail" for f in used_feats)
shots_n = sum(len(f["res"]["shots"]) for f in used_feats if f["res"])

cards = [
    (bn(total), "মোট ফিচার পরীক্ষা করা হয়েছে", BLUE_BG, NAVY),
    (bn(works), f"ঠিকমতো কাজ করছে ({pct(works, total)}%)", GREEN_BG, GREEN),
    (bn(partly), f"আংশিক দেখানো হয়েছে ({pct(partly, total)}%)", AMBER_BG, AMBER),
    (bn(broken), "কাজ করছে না", "FBE4E4" if broken else GREY_BG, "B01E1E" if broken else "555555"),
]
t = doc.add_table(rows=1, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
fixed_layout(t)
w = CONTENT_W / 4
for i, (num, label, bg, fg) in enumerate(cards):
    c = t.rows[0].cells[i]
    width(c, w)
    shade(c, bg)
    borders(c, "FFFFFF", 18)
    cell_margins(c, 160, 160, 120, 120)
    first_para(c, num, 30, True, fg, WD_ALIGN_PARAGRAPH.CENTER)
    p = para(c, label, 10, False, "333333", WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

para(doc, "", space_after=2)
# overall bar
t = doc.add_table(rows=1, cols=2 if partly else 1)
fixed_layout(t)
bar_w = CONTENT_W
gw = bar_w * works / total
cells = t.rows[0].cells
width(cells[0], gw if partly else bar_w)
shade(cells[0], GREEN)
first_para(cells[0], f"{pct(works, total)}% কাজ করছে", 10, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
cell_margins(cells[0], 70, 70)
if partly:
    width(cells[1], bar_w - gw)
    shade(cells[1], "E5A300")
    cell_margins(cells[1], 70, 70)
    first_para(cells[1], "", 8)
para(doc, f"সবুজ অংশ: ঠিকমতো কাজ করছে। হলুদ অংশ: আংশিক দেখানো হয়েছে। মোট {bn(shots_n)}টি পর্দার ছবি এই প্রতিবেদনে আছে।", 9.5, False, "555555", space_before=4, space_after=10)

# chapter table
heading(doc, "বিভাগ ধরে ফল", 2)
t = doc.add_table(rows=1, cols=5)
fixed_layout(t)
widths = [6.2, 1.8, 2.0, 2.0, 5.4]
for i, h in enumerate(["বিভাগ", "মোট", "কাজ করছে", "আংশিক", "অগ্রগতি"]):
    c = t.rows[0].cells[i]
    width(c, widths[i])
    shade(c, NAVY)
    cell_margins(c, 80, 80)
    first_para(c, h, 10, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER if i else None)
for k, i in enumerate(use):
    fs = [f for f in feats.values() if f["chapter"] == chapters[i - 1]]
    a = sum(f["status"] == "pass" for f in fs)
    b = sum(f["status"] == "skipped" for f in fs)
    row = t.add_row().cells
    bg = "FFFFFF" if k % 2 == 0 else "F7F9FC"
    for j in range(5):
        width(row[j], widths[j])
        shade(row[j], bg)
        borders(row[j], "E3E7ED", 4, ("bottom",))
        cell_margins(row[j], 70, 70)
    first_para(row[0], CH_BN[i][0], 10)
    first_para(row[1], bn(len(fs)), 10, False, None, WD_ALIGN_PARAGRAPH.CENTER)
    first_para(row[2], bn(a), 10, True, GREEN, WD_ALIGN_PARAGRAPH.CENTER)
    first_para(row[3], bn(b) if b else "-", 10, True if b else False, AMBER if b else "999999", WD_ALIGN_PARAGRAPH.CENTER)
    p0 = row[4].paragraphs[0]
    p0._element.getparent().remove(p0._element)
    inner = row[4].add_table(rows=1, cols=2 if b else 1)
    fixed_layout(inner)
    full = widths[4] - 0.5
    gwid = full * a / len(fs)
    width(inner.rows[0].cells[0], gwid if b else full)
    shade(inner.rows[0].cells[0], "2E9E5B")
    first_para(inner.rows[0].cells[0], "", 6)
    if b:
        width(inner.rows[0].cells[1], full - gwid)
        shade(inner.rows[0].cells[1], "E5A300")
        first_para(inner.rows[0].cells[1], "", 6)
    tail = row[4].add_paragraph()
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    ppr = tail._p.get_or_add_pPr()
    mark = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    mark.append(sz)
    ppr.append(mark)

# partly shown list
part_list = sorted((f for f in used_feats if f["status"] == "skipped"), key=lambda f: f["n"])
if part_list:
    heading(doc, "যা আংশিক দেখানো হয়েছে, এবং কেন", 2)
    para(doc, "নিচের বিষয়গুলোতে ফিচারটি তৈরি আছে এবং যতটুকু সম্ভব দেখানো হয়েছে। বাকি অংশ দেখানো যায়নি, কারণ সেটির জন্য বাইরের কোনো সেবার চাবি, সত্যিকারের অ্যাকাউন্ট বা বাস্তব পরিস্থিতি (যেমন ইন্টারনেট সত্যিই বন্ধ হওয়া) লাগে, অথবা অংশটির কাজ এখনও পুরো শেষ হয়নি। কারণ প্রতিটির নিচে লেখা আছে।", 10.5, False, "444444", space_after=6)
    for f in part_list:
        b = BN[f["n"]]
        t = doc.add_table(rows=1, cols=1)
        fixed_layout(t)
        c = t.rows[0].cells[0]
        width(c, CONTENT_W)
        shade(c, AMBER_BG)
        borders(c, "F0D58A", 4)
        cell_margins(c, 80, 80, 140, 140)
        p = first_para(c, f"{bn(f['n'])}. {b['name']}", 10.5, True, AMBER)
        para(c, b["N"], 10, False, "3A3A3A", space_after=0, space_before=2)
        para(doc, "", size=3, space_after=0)

# bugs found and fixed during testing
heading(doc, "পরীক্ষার সময় যে ত্রুটিগুলো ধরা পড়ে ঠিক করা হয়েছে", 2)
para(doc, "পরীক্ষা শুধু \"ঠিক আছে\" বলার জন্য নয়। পরীক্ষা চালাতে গিয়ে যেসব সত্যিকারের ভুল বেরিয়েছে, সেগুলো ঠিক করা হয়েছে এবং আবার পরীক্ষা করে দেখা হয়েছে। এর কয়েকটি:", 10.5, False, "444444", space_after=6)
for tx in (
    "ফিচার ৩২: শেষ হওয়া মেমোর মোট ও বাকির লাইনে ক্রেডিট আর টোকেনের হিসাব বাদ পড়ছিল। এখন সঠিক মোট দেখায়।",
    "ফিচার ২৩৪: হোলসেলার দর দেওয়ার পর পাতা নতুন করে না খোলা পর্যন্ত \"আপনার অফার\" দেখাত না। এখন সাথে সাথে বদলায়।",
    "ফিচার ২৩৭: এমপিওর মাল একবার স্টকে ঢোকার পরও আবার চাপলে দুবার ঢুকে যেত, আর দেরি গোনার ভুল হতো। এখন আটকানো হয়েছে।",
    "ফিচার ৩০৮: বিক্রেতার ক্যাশ অন ডেলিভারির সুইচ \"সেভ হয়েছে\" বলত কিন্তু আসলে সেভ হতো না। এখন সেভ হয়।",
    "ফিচার ৩১০: ফোনে \"কার্টে যোগ করুন\" চাপলে কোনো সাড়া দেখা যেত না। বার্তা এখন বাটনের ঠিক ওপরেই আসে।",
    "ফিচার ৩৫০: ফেরত বা বাতিল অর্ডারকেও ক্রেতাকে \"মাল পৌঁছালে টাকা দেবেন\" বলা হতো। এখন \"কিছু দিতে হবে না\" বলে।",
    "ফিচার ৩৭৬: নিষিদ্ধ শব্দ চিহ্নিত করার কাজ চলত, কিন্তু মালিকের শব্দের তালিকা লেখার কোনো পর্দা ছিল না। পর্দাটি যোগ করা হয়েছে।",
    "ফিচার ৪৫৫: নিজের সাইটের লিংকে পোর্ট নম্বর থাকলে মন্তব্যের কার্ড কোনো বার্তা ছাড়াই বাদ যেত। এখন ঠিক আছে।",
    "ফিচার ৫১৪: বাতিল হওয়া মেমোতেও সবুজ \"পরিশোধিত\" লাইন দেখাত। এখন লুকানো থাকে।",
    "ফিচার ৬১২: ফার্মাট্র্যাকের ডাটাবেস নতুন করে শুরু হলে বিক্রির নম্বর আবার প্রথম থেকে শুরু হতো, আর সাইট একটি বিক্রি নীরবে ফেলে দিত। আর ১২০ পিসের বক্সে গোনায় সামান্য ভগ্নাংশের গরমিল জমত। দুটিই ঠিক করা হয়েছে।",
):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(tx), 10.5)

# how to read
heading(doc, "প্রতিবেদনটি কীভাবে পড়বেন", 2)
for tx in (
    "প্রতিটি বিভাগের শুরুতে লেখা আছে ওই অংশে আপনি কী চেয়েছিলেন।",
    "প্রতিটি ফিচারের নিচে তিনটি অংশ আছে: ফিচারটি কী ও কেন লাগে, আমরা কীভাবে পরীক্ষা করেছি ও কী দেখেছি, আর পর্দার ছবি।",
    "সবুজ চিহ্ন: ফিচারটি যেভাবে চাওয়া হয়েছিল সেভাবে কাজ করছে।",
    "হলুদ চিহ্ন: ফিচারটি তৈরি আছে ও কাজ করছে, কিন্তু একটি অংশ দেখানো যায়নি। কারণটি পাশেই লেখা আছে।",
    "সব পরীক্ষা করা হয়েছে সাইটের একটি পরীক্ষার কপিতে, নমুনা তথ্য দিয়ে। আপনার আসল তথ্যে হাত দেওয়া হয়নি।",
    "প্রতিটি ছবির নিচে লেখা আছে ছবিটি কার চোখ দিয়ে দেখা (যেমন মালিক, একজন পাইকারি দোকান, একজন ক্রেতা)।",
):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(tx), 10.5)

# glossary
heading(doc, "কয়েকটি শব্দের মানে", 2)
t = doc.add_table(rows=0, cols=2)
fixed_layout(t)
for k, v in (
    ("পাইকারি দোকান", "ছোট দোকান বা ফার্মেসি, যারা আপনার কাছ থেকে মাঝারি পরিমাণে কেনে। খুচরা আর হোলসেলের মাঝামাঝি।"),
    ("হোলসেলার", "অনুমোদিত বড় ক্রেতা, আবার যারা নিজের পণ্য আপনার সাইটে তুলে বিক্রি করেন তারাও।"),
    ("এমপিও", "ওষুধ কোম্পানির বিক্রয় প্রতিনিধি।"),
    ("এমআরপি", "ওষুধের বক্সে ছাপা সর্বোচ্চ খুচরা মূল্য।"),
    ("বক্স, পাতা, পিস", "ওষুধের তিনটি মাপ। একটি বক্সে কয়েক পাতা থাকে, একটি পাতায় কয়েক পিস।"),
    ("মেমো ও ইনভয়েস", "বিক্রির রসিদ। ইনভয়েস হলো চূড়ান্ত বিল।"),
    ("ক্যাশ অন ডেলিভারি", "পণ্য হাতে পেয়ে টাকা দেওয়া।"),
    ("বিকাশ, নগদ", "মোবাইলে টাকা পাঠানোর সেবা।"),
):
    row = t.add_row().cells
    width(row[0], 4.4)
    width(row[1], CONTENT_W - 4.4)
    for c in row:
        borders(c, "E3E7ED", 4, ("bottom",))
        cell_margins(c, 60, 60)
    first_para(row[0], k, 10, True, NAVY)
    first_para(row[1], v, 10)

# ------------------------------------------------------------------ chapters
for i in use:
    doc.add_page_break()
    title, intro = CH_BN[i]
    heading(doc, f"{bn(i)}. {title}", 1)
    t = doc.add_table(rows=1, cols=1)
    fixed_layout(t)
    c = t.rows[0].cells[0]
    width(c, CONTENT_W)
    shade(c, BLUE_BG)
    borders(c, "C9D6EA", 4)
    cell_margins(c, 120, 120, 180, 180)
    first_para(c, "এই অংশে আপনি যা চেয়েছিলেন", 10.5, True, NAVY)
    para(c, intro, 10.5, False, "222222", space_after=0, space_before=3)
    para(doc, "", size=4, space_after=2)

    for f in sorted((x for x in feats.values() if x["chapter"] == chapters[i - 1]), key=lambda x: x["n"]):
        b = BN[f["n"]]
        ok = f["status"] == "pass"
        label = "ঠিকমতো কাজ করছে" if ok else "আংশিক দেখানো হয়েছে"
        fg, bg = (GREEN, GREEN_BG) if ok else (AMBER, AMBER_BG)

        t = doc.add_table(rows=1, cols=2)
        fixed_layout(t)
        a, z = t.rows[0].cells
        width(a, CONTENT_W - 4.6)
        width(z, 4.6)
        shade(a, GREY_BG)
        shade(z, bg)
        for c in (a, z):
            borders(c, "FFFFFF", 4)
            cell_margins(c, 90, 90, 140, 140)
        first_para(a, f"{bn(f['n'])}. {b['name']}", 12.5, True, NAVY)
        a.paragraphs[0].paragraph_format.keep_with_next = True
        first_para(z, label, 10, True, fg, WD_ALIGN_PARAGRAPH.CENTER)
        tr = t.rows[0]._tr
        trpr = tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trpr.append(cs)

        p = para(doc, "", space_before=6, space_after=2, keep=True)
        set_font(p.add_run("ফিচারটি কী ও কেন লাগে"), 10.5, True, "444444")
        para(doc, b["E"], 11, space_after=5)
        p = para(doc, "", space_after=2, keep=True)
        set_font(p.add_run("আমরা কীভাবে পরীক্ষা করেছি ও কী দেখেছি"), 10.5, True, "444444")
        para(doc, b["T"], 11, space_after=5)
        if b["N"]:
            t = doc.add_table(rows=1, cols=1)
            fixed_layout(t)
            c = t.rows[0].cells[0]
            width(c, CONTENT_W)
            shade(c, AMBER_BG)
            borders(c, "F0D58A", 4)
            cell_margins(c, 80, 80, 140, 140)
            first_para(c, "জেনে রাখা ভালো", 10, True, AMBER)
            para(c, b["N"], 10, False, "3A3A3A", space_after=0, space_before=2)
            para(doc, "", size=3, space_after=2)

        shots = f["res"]["shots"] if f["res"] else []
        for k in range(0, len(shots), 2):
            pair = shots[k:k + 2]
            t = doc.add_table(rows=2, cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            fixed_layout(t)
            for j in range(2):
                cap, pic = t.rows[0].cells[j], t.rows[1].cells[j]
                width(cap, CONTENT_W / 2)
                width(pic, CONTENT_W / 2)
                for cc in (cap, pic):
                    cell_margins(cc, 50, 50, 70, 70)
                if j < len(pair):
                    s = pair[j]
                    shade(cap, "EEF1F5")
                    first_para(cap, f"ধাপ {bn(k + j + 1)}: {b['S'][k + j]}", 9, True, "1F2937")
                    cap.paragraphs[0].paragraph_format.keep_with_next = True
                    pp = pic.paragraphs[0]
                    pp.paragraph_format.keep_with_next = True
                    pp.add_run().add_picture(str(thumb(s["file"])), width=Cm(8.2))
                    p2 = pic.add_paragraph()
                    p2.paragraph_format.space_after = Pt(0)
                    set_font(p2.add_run("যাঁর চোখে দেখা: " + ROLE_BN.get(s["role"], "একজন ব্যবহারকারী")), 8, None, "6B7280")
            trpr = t.rows[0]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            t.rows[1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        para(doc, "", size=6, space_after=6)

doc.save(OUT)
print(f"saved {OUT} ({OUT.stat().st_size / 1e6:.1f} MB): {total} features, chapters {use}")
